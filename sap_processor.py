import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import sys
import os

# --- KONFIGURACE ---
REQUIRED_COLS = {
    "material": "Material",
    "material description": "Material description",
    "batch": "Batch",
    "total quantity": "Total Quantity",
    "storage location": "Storage location"
}

OPTIONAL_COLS = {
    "plant": "Plant"
}

# --- FUNKCE ---

def normalize_columns(df):
    """
    Přejmenuje sloupce v DataFrame na standardní formát (dle zadání),
    aby nám nezáleželo na velikosti písmen v originále.
    """
    actual_cols_lower = {c.lower().strip(): c for c in df.columns}
    
    mapping = {}
    missing = []
    
    for req_key, req_std_name in REQUIRED_COLS.items():
        if req_key in actual_cols_lower:
            mapping[actual_cols_lower[req_key]] = req_std_name
        else:
            missing.append(req_std_name)
            
    for opt_key, opt_std_name in OPTIONAL_COLS.items():
        if opt_key in actual_cols_lower:
            mapping[actual_cols_lower[opt_key]] = opt_std_name

    if missing:
        raise ValueError(f"Chybí povinné sloupce: {', '.join(missing)}")
        
    df = df.rename(columns=mapping)
    return df

def find_best_sheet(file_path):
    """
    Najde list s největší datovou plochou.
    """
    xl = pd.ExcelFile(file_path)
    best_sheet = None
    max_area = -1
    
    for sheet_name in xl.sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet_name)
        area = df.shape[0] * df.shape[1]
        
        if area > max_area and not df.empty:
            max_area = area
            best_sheet = sheet_name
            
    if best_sheet is None:
        raise ValueError("Nenašel jsem žádný list s daty.")
        
    print(f"Vybrán list pro zpracování: '{best_sheet}' (Plocha: {max_area} buněk)")
    return best_sheet

def process_file(input_path, output_path):
    """
    Hlavní logika zpracování jednoho souboru.
    """
    print(f"--- Zpracovávám soubor: {input_path} ---")
    
    try:
        # 1. Autodetekce listu
        sheet_name = find_best_sheet(input_path)
        df = pd.read_excel(input_path, sheet_name=sheet_name)
        
        # 2. Normalizace a Validace sloupců
        df = normalize_columns(df)
        
        # 3. Filtr Storage location (F010, F070)
        allowed_locations = ["F010", "F070"]
        df = df[df["Storage location"].astype(str).isin(allowed_locations)].copy()
        
        # 4. Smazání sloupců
        cols_to_drop = ["Storage location"]
        if "Plant" in df.columns:
            cols_to_drop.append("Plant")
        df = df.drop(columns=cols_to_drop, errors='ignore')
        
        # 5. Uspořádání sloupců
        ordered_cols = ["Material", "Material description", "Batch", "Total Quantity"]
        remaining_cols = [c for c in df.columns if c not in ordered_cols]
        df = df[ordered_cols + remaining_cols]
        
        # 6. Přejmenování
        rename_map = {
            "Material description": "Název",
            "Total Quantity": "Mnozstvi_SAP"
        }
        df = df.rename(columns=rename_map)
        
        # 7. Sort
        df["Mnozstvi_SAP"] = pd.to_numeric(df["Mnozstvi_SAP"], errors='coerce').fillna(0)
        df = df.sort_values(by="Mnozstvi_SAP", ascending=False, kind='mergesort')
        
        # 8. Smazat první datový řádek po sortu
        if len(df) > 0:
            df = df.iloc[1:]
        
        # 9. Ořezat na 4 sloupce
        df = df.iloc[:, :4]
        
        # --- EXPORT DO EXCELU S TABULKOU ---
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='SAP', index=False)
            
        wb = openpyxl.load_workbook(output_path)
        ws = wb["SAP"]
        
        max_row = ws.max_row
        ref = f"A1:D{max_row}"
        
        tab = Table(displayName="tbl_SAP", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        
        ws.add_table(tab)
        wb.save(output_path)
        
        print(f"✅ Hotovo. Uloženo do: {output_path}")

    except Exception as e:
        print(f"❌ Chyba při zpracování: {e}")
        sys.exit(1)

# --- SPUŠTĚNÍ ---
if __name__ == "__main__":
    # Pevně definované cesty
    input_dir = "sklady_porovnani/input"
    output_dir = "sklady_porovnani/output"
    filename = "SAP.xlsx"
    
    infile = os.path.join(input_dir, filename)
    outfile = os.path.join(output_dir, filename)
    
    # Zajištění výstupní složky
    os.makedirs(output_dir, exist_ok=True)

    # Kontrola, zda existuje vstup
    if not os.path.exists(infile):
        print(f"❌ CHYBA: Soubor '{infile}' neexistuje.")
        print("Nahraj prosím soubor 'SAP.xlsx' do složky 'sklady_porovnani/input/'.")
        sys.exit(1)

    # Spuštění
    process_file(infile, outfile)