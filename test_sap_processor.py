import unittest
import pandas as pd
import openpyxl
import os
from sap_processor import process_file

class TestSapProcessor(unittest.TestCase):
    
    def setUp(self):
        """
        Příprava před každým testem: Vytvoříme testovací Excel s 'bordel' daty.
        """
        self.input_file = "test_input_temp.xlsx"
        self.output_file = "test_output_temp.xlsx"
        
        # Vytvoříme data, která testují všechny naše podmínky
        data = {
            # Test case-insensitivity (různá velikost písmen)
            "mATERIAL": ["M1", "M2", "M3", "M4"],
            "Material description": ["Desc A", "Desc B", "Desc C", "Desc D"],
            "baTCH": ["B1", "B2", "B3", "B4"],
            # Test numerického řazení (čísla jako stringy, NaN)
            "Total Quantity": ["10", "100", "5", "chyba"], 
            "Storage location": ["F010", "F070", "X999", "F010"], # X999 by se mělo smazat
            "Plant": ["P1", "P1", "P1", "P1"], # Tento sloupec se má smazat
            "NecoNavic": ["X", "X", "X", "X"] # Tento sloupec se má taky oříznout
        }
        
        df = pd.DataFrame(data)
        
        # Uložíme jako Excel
        df.to_excel(self.input_file, sheet_name="Sheet1", index=False)
        print(f"\n[SETUP] Vytvořen testovací soubor: {self.input_file}")

    def test_process_logic(self):
        """
        Hlavní testovací logika.
        """
        # 1. Spustíme naši funkci
        print("[TEST] Spouštím process_file...")
        process_file(self.input_file, self.output_file)
        
        # 2. Ověříme, že výstupní soubor existuje
        self.assertTrue(os.path.exists(self.output_file), "Výstupní soubor nebyl vytvořen.")
        
        # 3. Načteme výsledek pro kontrolu dat (pomocí pandas)
        df_out = pd.read_excel(self.output_file, sheet_name="SAP")
        
        # KONTROLA A: Správné sloupce a jejich pořadí
        expected_cols = ["Material", "Název", "Batch", "Mnozstvi_SAP"]
        self.assertListEqual(list(df_out.columns), expected_cols, "Názvy nebo pořadí sloupců nesedí.")
        
        # KONTROLA B: Filtrace (Storage location)
        # Měli jsme 4 řádky: F010, F070, X999 (smazat), F010. 
        # Zbyly by 3. ALE logika říká "Smaž řádek 2 (první datový) po sortu".
        # Takže očekáváme 2 řádky.
        self.assertEqual(len(df_out), 2, "Počet řádků neodpovídá (filtr + smazání prvního řádku).")
        
        # KONTROLA C: Řazení (Sestupně dle Mnozstvi_SAP)
        # Původní data po filtraci (F010/F070):
        # M1: 10
        # M2: 100
        # M4: chyba -> 0 (fillna)
        # Seřazeno: 100 (M2), 10 (M1), 0 (M4).
        # První řádek (M2, 100) se má smazat.
        # Zbývá: M1 (10) a M4 (0).
        # První řádek ve výsledku by měl mít Mnozstvi_SAP = 10.
        self.assertEqual(df_out.iloc[0]["Mnozstvi_SAP"], 10, "Řazení nebo smazání prvního řádku neproběhlo správně.")
        
        # 4. Ověříme formátování Excelu (pomocí openpyxl)
        wb = openpyxl.load_workbook(self.output_file)
        ws = wb["SAP"]
        
        # KONTROLA D: Existence tabulky "tbl_SAP"
        self.assertIn("tbl_SAP", ws.tables, "Tabulka 'tbl_SAP' nebyla v Excelu nalezena.")
        
        # KONTROLA E: Rozsah tabulky
        table_range = ws.tables["tbl_SAP"].ref
        # Máme hlavičku + 2 řádky dat = 3 řádky celkem. Rozsah A1:D3
        self.assertEqual(table_range, "A1:D3", f"Rozsah tabulky je špatně: {table_range}")
        
        print("[TEST] Všechny kontroly prošly úspěšně!")

    def tearDown(self):
        """
        Úklid po testu: Smažeme dočasné soubory.
        """
        if os.path.exists(self.input_file):
            os.remove(self.input_file)
        if os.path.exists(self.output_file):
            os.remove(self.output_file)
        print("[TEARDOWN] Dočasné soubory smazány.")

if __name__ == "__main__":
    unittest.main()