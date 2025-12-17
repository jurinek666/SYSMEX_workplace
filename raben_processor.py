import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
import sys
import os

def process_raben(input_path, output_path):
    print(f"--- Zpracovávám RABEN: {input_path} ---")
    
    try:
        # 1. Načtení dat (první list)
        df = pd.read_excel(input_path, sheet_name=0)
        
        # 2. Mapování sloupců
        # Normalizujeme názvy vstupních sloupců (strip whitespace)
        df.columns = [str(c).strip() for c in df.columns]
        
        col_map = {
            "1-Císlo zboží": "Material",
            "3-název": "Nazev",
            "12-šarže": "Batch",
            "4-ks": "Mnozstvi_RABEN"
        }
        
        # Kontrola existence sloupců
        missing_cols = [c for c in col_map.keys() if c not in df.columns]
        if missing_cols:
            raise ValueError(f"Chybí sloupce v RABEN souboru: {missing_cols}")
            
        # Přejmenování a výběr
        df = df.rename(columns=col_map)
        df = df[list(col_map.values())] # Ponechat jen tyto 4
        
        # 3. Datové typy
        # Textové sloupce
        for col in ["Material", "Nazev", "Batch"]:
            df[col] = df[col].astype(str).replace("nan", "").str.strip()
            
        # Numerický sloupec
        df["Mnozstvi_RABEN"] = pd.to_numeric(df["Mnozstvi_RABEN"], errors='coerce').fillna(0)
        
        # 4. Export
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='RABEN', index=False)
            
        # 5. Formátování tabulky
        wb = openpyxl.load_workbook(output_path)
        ws = wb["RABEN"]
        max_row = ws.max_row
        ref = f"A1:D{max_row}"
        
        tab = Table(displayName="tbl_RABEN", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)
        
        wb.save(output_path)
        print(f"✅ RABEN hotovo. Uloženo do: {output_path}")

    except Exception as e:
        print(f"❌ Chyba RABEN: {e}")
        sys.exit(1)

if __name__ == "__main__":
    input_dir = "sklady_porovnani/input"
    output_dir = "sklady_porovnani/output"
    filename = "RABEN.xlsx"
    
    infile = os.path.join(input_dir, filename)
    outfile = os.path.join(output_dir, filename)
    
    os.makedirs(output_dir, exist_ok=True)
    
    if not os.path.exists(infile):
        print(f"❌ Soubor {infile} neexistuje.")
        sys.exit(1)
        
    process_raben(infile, outfile)